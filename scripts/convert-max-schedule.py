"""Convert the MAX "Staff Tracker & Booth Schedule" xlsx (one tab per area,
day/time rows x demo-station columns) into the flat schedule schema the app reads.

Usage: python convert-max-schedule.py <input.xlsx> <output.json>
"""
import datetime
import difflib
import json
import re
import sys

import openpyxl

ROSTER_TAB = "Adobe Booth Approved Staff"
DAY_DATES = {"Tuesday": "2026-11-10", "Wednesday": "2026-11-11", "Thursday": "2026-11-12"}
SLOT_MINUTES = 30
NOT_A_PERSON = {"-", "break", "lunch", "demo", "demo prep", "demo presentation",
                "community expert", "community expert?", "float", "on"}
ALIASES = {"Jusin Patton": "Justin Patton"}


def norm(s):
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ").replace("​", "")).strip()


def load_roster(wb):
    ws = wb[ROSTER_TAB]
    people = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        area, first, last = (norm(v) if v else "" for v in r[1:4])
        if first and last:
            people.append({"area": area, "first": first, "full": f"{first} {last}"})
    return people


def area_matches(roster_area, tab):
    a, t = roster_area.lower(), tab.lower()
    return t in a or a.replace("adobe ", "") in t


class Resolver:
    def __init__(self, roster):
        self.roster = roster
        self.full_names = {p["full"].lower(): p["full"] for p in roster}
        self.log = {}

    def resolve(self, raw, tab, tab_full_names):
        name = ALIASES.get(raw, raw)
        if name.lower() in self.full_names:
            return self.full_names[name.lower()]
        in_area = [p for p in self.roster if area_matches(p["area"], tab)]
        if " " not in name:
            in_tab = [{"first": n.split(" ")[0], "full": n} for n in tab_full_names]
            for pool in (in_tab, in_area, self.roster):
                hits = {p["full"] for p in pool if p["first"].lower() == name.lower()}
                if len(hits) == 1:
                    full = hits.pop()
                    return self._note(raw, self.resolve(full, tab, ()) if pool is in_tab else full, tab)
            hits = {p["full"] for p in in_area if p["first"].lower().startswith(name.lower())}
            if len(hits) == 1:
                return self._note(raw, hits.pop(), tab)
            return self._note(raw, name, tab, unresolved=True)
        for pool in (in_area, self.roster):
            close = difflib.get_close_matches(name, [p["full"] for p in pool], n=1, cutoff=0.85)
            if close:
                return self._note(raw, close[0], tab)
        return self._note(raw, name, tab, unresolved=True)

    def _note(self, raw, out, tab, unresolved=False):
        if raw != out or unresolved:
            self.log[(tab, raw)] = (out, unresolved)
        return out


def split_people(cell):
    """'A, B' / 'A + B' -> ['A','B']; drops '- Expert' suffix and break/lunch/role notes."""
    out, skipped = [], []
    for part in re.split(r",| \+ ", norm(cell)):
        p = re.sub(r"\s*-\s*Expert$", "", part.strip(), flags=re.I)
        if not p:
            continue
        if p.lower() in NOT_A_PERSON or "(" in p or ":" in p:
            skipped.append(p)
            continue
        out.append(p)
    return out, skipped


def station_columns(ws):
    """Header cells from column C up to and including FLOATER (or the first gap)."""
    cols = {}
    for cell in ws[2][2:]:
        h = norm(cell.value) if cell.value else ""
        if not h:
            break
        cols[cell.column] = h
        if h.upper() == "FLOATER":
            break
    return cols


def label(header):
    h = header.title() if header.isupper() or header.upper().startswith("DEMO") else header
    return re.sub(r"\bAbi\b", "ABI", re.sub(r"\bFfcpe\b", "FFCPE", h))


def read_slots(ws, tab, warnings):
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        day, t = (norm(r[0]) if r[0] else ""), r[1]
        if day in DAY_DATES and isinstance(t, datetime.time):
            rows.append([day, t, r])
    for i in range(1, len(rows) - 1):
        prev, cur, nxt = rows[i - 1], rows[i], rows[i + 1]
        if cur[0] != prev[0] and nxt[0] == cur[0] and nxt[1] < cur[1]:
            warnings.append(f"{tab}: row '{cur[0]} {cur[1]:%H:%M}' treated as {prev[0]} (out of sequence)")
            cur[0] = prev[0]
    return rows


def add_minutes(t, m):
    return (datetime.datetime.combine(datetime.date(2000, 1, 1), t) + datetime.timedelta(minutes=m)).strftime("%H:%M")


def convert(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    resolver = Resolver(load_roster(wb))
    warnings, skipped_notes, out = [], {}, []

    for ws in wb.worksheets:
        if ws.title == ROSTER_TAB:
            continue
        tab = norm(ws.title)
        cols = station_columns(ws)
        slots = read_slots(ws, tab, warnings)
        tab_full_names = {p for _, _, r in slots for col in cols
                          if col - 1 < len(r) and isinstance(r[col - 1], str)
                          for p in split_people(r[col - 1])[0] if " " in p}
        runs = {}  # (person, station) -> open run
        for day, t, r in slots:
            seen_now = set()
            for col, header in cols.items():
                val = r[col - 1] if col - 1 < len(r) else None
                if not isinstance(val, str) or not norm(val):
                    continue
                people, skipped = split_people(val)
                for s in skipped:
                    skipped_notes[s] = skipped_notes.get(s, 0) + 1
                station = f"{tab} · {label(header)}"
                for raw in people:
                    person = resolver.resolve(raw, tab, tab_full_names)
                    seen_now.add((person, station))
                    key = (person, station)
                    run = runs.get(key)
                    start = t.strftime("%H:%M")
                    if run and run["day"] == day and run["end_time"] == start:
                        run["end_time"] = add_minutes(t, SLOT_MINUTES)
                    else:
                        if run:
                            out.append(run)
                        runs[key] = {"staff_name": person, "day": day, "date": DAY_DATES[day],
                                     "start_time": start, "end_time": add_minutes(t, SLOT_MINUTES),
                                     "station": station, "is_floater": header.upper() == "FLOATER",
                                     "group": tab}
        out.extend(runs.values())

    out.sort(key=lambda s: (s["date"], s["group"], s["station"], s["start_time"], s["staff_name"]))
    return out, warnings, resolver.log, skipped_notes


def double_bookings(rows):
    by_person = {}
    for r in rows:
        by_person.setdefault((r["staff_name"], r["date"]), []).append(r)
    clashes = []
    for (name, date), shifts in by_person.items():
        shifts.sort(key=lambda s: s["start_time"])
        for a, b in zip(shifts, shifts[1:]):
            if b["start_time"] < a["end_time"]:
                clashes.append(f"{name} {date}: {a['station']} {a['start_time']}-{a['end_time']} "
                               f"overlaps {b['station']} {b['start_time']}-{b['end_time']}")
    return clashes


if __name__ == "__main__":
    in_path, out_path = sys.argv[1], sys.argv[2]
    rows, warnings, name_log, skipped = convert(in_path)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=2, ensure_ascii=False)
    print(f"Wrote {len(rows)} shifts, {len({r['staff_name'] for r in rows})} people, "
          f"{len({r['station'] for r in rows})} stations, {len({r['group'] for r in rows})} areas -> {out_path}")
    for w in warnings:
        print("WARNING:", w)
    print("\nNAME RESOLUTIONS:")
    for (tab, raw), (out, unresolved) in sorted(name_log.items()):
        print(f"  [{tab}] {raw!r} -> {out!r}" + ("   (UNRESOLVED, kept as-is)" if unresolved else ""))
    print("\nSKIPPED CELL TEXT (not shifts):")
    for s, n in sorted(skipped.items()):
        print(f"  {n:3d} x {s!r}")
    clashes = double_bookings(rows)
    print(f"\nDOUBLE BOOKINGS: {len(clashes)}")
    for c in clashes:
        print("  " + c)
