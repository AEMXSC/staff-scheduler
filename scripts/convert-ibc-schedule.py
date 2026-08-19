"""Convert the IBC grid-format staff xlsx (person rows x time-slot columns)
into the flat schedule schema the app reads.

Usage: python convert-ibc-schedule.py <input.xlsx> <output.json>
"""
import json
import re
import sys

import openpyxl

SECTION_HEADERS = {"Demo Artists", "ATE Floaters", "Info Counter", "Partner Booth"}
SHEET_DATES = {
    "Friday": "2026-09-11",
    "Saturday": "2026-09-12",
    "Sunday": "2026-09-13",
    "Monday": "2026-09-14",
}


def parse_time(label):
    """'2:30 P.M.' -> '14:30' (24h). Minutes are optional ('4 P.M.' -> '16:00')."""
    m = re.match(r"(\d{1,2})(?::(\d{2}))?\s*([AP])\.?M\.?", label.strip(), re.I)
    if not m:
        return None
    h, mnt, ap = int(m.group(1)), int(m.group(2) or 0), m.group(3).upper()
    if ap == "P" and h != 12:
        h += 12
    if ap == "A" and h == 12:
        h = 0
    return f"{h:02d}:{mnt:02d}"


def shift_end_times(ws, sheet_name):
    """Row 1/2 hold 'Shift N - START - END' labels; return [end1, end2] as HH:MM."""
    ends = []
    for r in (1, 2):
        label = ws.cell(r, 2).value or ""
        m = re.search(r"[-–—]\s*([\d:APM\.\s]+)$", label.strip())
        end = parse_time(m.group(1)) if m else None
        if end is None:
            print(f"WARNING: {sheet_name}: could not parse shift {r} end time from {label!r}")
        ends.append(end)
    return ends


def build_column_map(ws, sheet_name, num_shifts=2):
    """Map each column index (1-based) with a time header to (start_time, shift_idx).

    Shift 1 and shift 2 columns are separated by exactly one blank spacer
    column in the sheet layout -- use that as the boundary marker rather than
    comparing time values, since shift 2's first label can equal or precede
    shift 1's last label (e.g. both printed as "1:30 P.M.").
    """
    header_row = 5
    cols = {}
    shift = 0
    pending_bump = False
    for cell in ws[header_row]:
        val = cell.value
        if not val or val == "\xa0":
            if cols:
                pending_bump = True
            continue
        if pending_bump:
            shift += 1
            pending_bump = False
        if shift > num_shifts - 1:
            raise ValueError(
                f"{sheet_name}: found more than {num_shifts} shift blocks in the header "
                f"row (an extra blank spacer column?) -- update shift_end_times' row range "
                f"to match, or check column {cell.coordinate} manually."
            )
        t = parse_time(str(val))
        if not t:
            continue
        cols[cell.column] = {"start": t, "shift": shift}
    return cols


def col_end_time(col_indices, idx_pos, cols, shift_ends):
    """End time for the slot at cols position idx_pos in the sorted column list."""
    cur_col = col_indices[idx_pos]
    cur = cols[cur_col]
    if idx_pos + 1 < len(col_indices):
        nxt = cols[col_indices[idx_pos + 1]]
        if nxt["shift"] == cur["shift"]:
            return nxt["start"]
    return shift_ends[cur["shift"]]


def convert(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows_out = []
    warnings = []

    for sheet_name in wb.sheetnames:
        if sheet_name not in SHEET_DATES:
            continue
        ws = wb[sheet_name]
        cols = build_column_map(ws, sheet_name)
        col_indices = sorted(cols)
        shift_ends = shift_end_times(ws, sheet_name)
        date = SHEET_DATES[sheet_name]

        current_group = None
        for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
            name_cell = row[0].value
            name = name_cell.strip() if isinstance(name_cell, str) else None
            if not name or name == "\xa0":
                continue

            cell_by_col = {c.column: c.value for c in row}
            has_any_slot = any(
                cell_by_col.get(col) not in (None, "\xa0") for col in col_indices
            )

            if name in SECTION_HEADERS:
                current_group = name
                continue

            if not has_any_slot:
                continue  # unfilled placeholder row (e.g. "Frame person 6")

            if "TBD" in name.upper():
                warnings.append(f"{sheet_name}: unresolved TBD row -> {name!r}")

            # Walk columns left-to-right, collapsing contiguous equal stations.
            run_station = None
            run_start_idx = None
            for i, col in enumerate(col_indices):
                val = cell_by_col.get(col)
                station = val.strip() if isinstance(val, str) and val != "\xa0" else None

                if station != run_station:
                    if run_station is not None:
                        rows_out.append({
                            "staff_name": name,
                            "day": sheet_name,
                            "date": date,
                            "start_time": cols[col_indices[run_start_idx]]["start"],
                            "end_time": col_end_time(col_indices, i - 1, cols, shift_ends),
                            "station": run_station,
                            "is_floater": "floater" in run_station.lower(),
                            "group": current_group,
                        })
                    run_station = station
                    run_start_idx = i

            if run_station is not None:
                rows_out.append({
                    "staff_name": name,
                    "day": sheet_name,
                    "date": date,
                    "start_time": cols[col_indices[run_start_idx]]["start"],
                    "end_time": col_end_time(col_indices, len(col_indices) - 1, cols, shift_ends),
                    "station": run_station,
                    "is_floater": "floater" in run_station.lower(),
                    "group": current_group,
                })

    return rows_out, warnings


if __name__ == "__main__":
    in_path, out_path = sys.argv[1], sys.argv[2]
    rows, warnings = convert(in_path)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=2)
    print(f"Wrote {len(rows)} rows to {out_path}")
    for w in warnings:
        print("WARNING:", w)
